from datetime import datetime
from io import BytesIO
import re
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ILLEGAL_EXCEL_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")
NAVY = "152238"
GOLD = "D6A84B"
LIGHT = "F4F6F8"


CATALOG_COLUMNS = [
    ("agency", "Órgão"),
    ("object", "Objeto"),
    ("city", "Município"),
    ("state", "UF"),
    ("modality", "Modalidade"),
    ("nature", "Tipo"),
    ("estimated_value", "Valor estimado"),
    ("srp", "Registro de preços"),
    ("published_at", "Publicado em"),
    ("opening_at", "Início das propostas"),
    ("closing_at", "Fim das propostas"),
    ("source_name", "Portal de disputa"),
    ("source_channel", "Fonte do LicitaNexo"),
    ("pncp_control_number", "Nº PNCP"),
    ("source_url", "Link oficial"),
]

SAVED_COLUMNS = [
    ("agency", "Órgão"),
    ("object", "Objeto"),
    ("city", "Município"),
    ("state", "UF"),
    ("modality", "Modalidade"),
    ("stage_display", "Status"),
    ("certame_at", "Data do certame"),
    ("estimated_value", "Valor estimado"),
    ("source_name", "Portal de disputa"),
    ("source_channel", "Fonte do LicitaNexo"),
    ("pncp_control_number", "Nº PNCP"),
    ("notes", "Minhas anotações"),
    ("source_url", "Link oficial"),
]


def _safe_text(value):
    if value is None:
        return ""
    return ILLEGAL_EXCEL_CHARS.sub(" ", str(value))


def _clean_frame(rows):
    frame = pd.DataFrame(rows)
    for column in frame.columns:
        if frame[column].dtype == "object":
            frame[column] = frame[column].map(_safe_text)
    return frame


def _format_date(value):
    if not value:
        return ""
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        return parsed.strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        text = str(value).replace("T", " ")
        return text[:16]


def _format_brl(value):
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        return ""
    if number == 0:
        return ""
    return f"R$ {number:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _yes_no(value):
    if value in (1, True, "1", "true", "True", "SIM", "Sim", "sim"):
        return "Sim"
    if value in (0, False, "0", "false", "False", "NAO", "Não", "Nao", "não", "nao"):
        return "Não"
    return _safe_text(value)


def _translated_rows(rows, columns):
    translated = []
    for row in rows:
        item = {}
        for key, label in columns:
            value = row.get(key)
            if key == "estimated_value":
                value = _format_brl(value)
            elif key in {"published_at", "opening_at", "closing_at", "certame_at"}:
                value = _format_date(value)
            elif key == "srp":
                value = _yes_no(value)
            item[label] = _safe_text(value)
        translated.append(item)
    return translated


def _style_workbook(writer, data_sheet_name="Editais"):
    for sheet in writer.book.worksheets:
        sheet.sheet_view.showGridLines = False
        # Só a planilha de dados precisa de congelamento/filtro.
        # Congelar e depois descongelar a aba Resumo fazia o openpyxl deixar
        # <selection pane="bottomLeft"> sem um <pane> correspondente no XML;
        # o Microsoft Excel então reparava /xl/worksheets/sheet1.xml ao abrir.
        if sheet.title == data_sheet_name:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        else:
            sheet.auto_filter.ref = None
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 28
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx, column in enumerate(sheet.columns, 1):
            values = [len(str(c.value or "")) for c in column]
            width = min(max(max(values, default=8) + 2, 12), 55)
            header = str(sheet.cell(1, idx).value or "")
            if header == "Objeto":
                width = 55
            elif header in {"Link oficial", "Minhas anotações"}:
                width = 42
            elif header in {"UF", "Status"}:
                width = max(width, 12)
            sheet.column_dimensions[get_column_letter(idx)].width = width

        if sheet.title == data_sheet_name and sheet.max_row >= 2:
            headers = {sheet.cell(1, col).value: col for col in range(1, sheet.max_column + 1)}
            link_col = headers.get("Link oficial")
            if link_col:
                for row in range(2, sheet.max_row + 1):
                    cell = sheet.cell(row, link_col)
                    url = str(cell.value or "").strip()
                    if url.startswith(("http://", "https://")):
                        # Compatibilidade máxima com Microsoft Excel:
                        # não cria fórmula HYPERLINK nem relacionamento externo no XLSX.
                        # A URL permanece como texto puro, evitando reparos em sheet XML.
                        clean_url = ILLEGAL_EXCEL_CHARS.sub("", url).strip()
                        cell.value = clean_url[:32767]
                        cell.number_format = "@"

    if "Resumo" in writer.book.sheetnames:
        summary = writer.book["Resumo"]
        summary.sheet_view.showGridLines = False
        for cell in summary[1]:
            cell.fill = PatternFill("solid", fgColor=GOLD)
            cell.font = Font(color=NAVY, bold=True)


def pipeline_excel(bundle, company_name):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        summary = pd.DataFrame([{
            "Empresa": _safe_text(company_name),
            "Gerado em": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "Oportunidades": len(bundle.get("Pipeline", [])),
            "Tarefas": len(bundle.get("Tarefas", [])),
            "Documentos": len(bundle.get("Documentos", [])),
            "Itens de preço": len(bundle.get("Precos", [])),
        }])
        summary.to_excel(writer, sheet_name="Resumo", index=False)
        for sheet, rows in bundle.items():
            _clean_frame(rows).to_excel(writer, sheet_name=sheet[:31], index=False)
        _style_workbook(writer, "Pipeline")
    return output.getvalue()


def catalog_excel(rows, company_name, filters_text=""):
    output = BytesIO()
    translated = _translated_rows(rows, CATALOG_COLUMNS)
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([{
            "Empresa": _safe_text(company_name),
            "Pesquisa / filtros": _safe_text(filters_text or "Nenhum"),
            "Gerado em": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "Oportunidades encontradas": len(rows),
        }]).to_excel(writer, sheet_name="Resumo", index=False)
        pd.DataFrame(translated).to_excel(writer, sheet_name="Editais", index=False)
        _style_workbook(writer, "Editais")
    return output.getvalue()


def saved_editals_excel(rows, company_name):
    output = BytesIO()
    translated = _translated_rows(rows, SAVED_COLUMNS)
    status_counts = {}
    for row in rows:
        status = row.get("stage_display") or "Salvo"
        status_counts[status] = status_counts.get(status, 0) + 1
    summary_data = {
        "Empresa": _safe_text(company_name),
        "Gerado em": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "Total em Meus Editais": len(rows),
    }
    for status, count in sorted(status_counts.items()):
        summary_data[status] = count
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame([summary_data]).to_excel(writer, sheet_name="Resumo", index=False)
        pd.DataFrame(translated).to_excel(writer, sheet_name="Meus Editais", index=False)
        _style_workbook(writer, "Meus Editais")
    return output.getvalue()


def _pdf_header(story, title, company_name, styles):
    story.append(Paragraph("<b>LicitaNexo</b>", styles["Title"]))
    story.append(Paragraph(title, styles["Heading2"]))
    story.append(Paragraph(
        f"Empresa: {escape(_safe_text(company_name))} · Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}",
        styles["BodyText"],
    ))
    story.append(Spacer(1, 5 * mm))


def _pdf_table(data, widths):
    table = Table(data, repeatRows=1, colWidths=widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#152238")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F8")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def pipeline_pdf(rows, company_name):
    return saved_editals_pdf(rows, company_name)


def catalog_pdf(rows, company_name, filters_text=""):
    output = BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10 * mm,
                            leftMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    story = []
    _pdf_header(story, "Editais pesquisados", company_name, styles)
    story.append(Paragraph(
        f"Pesquisa / filtros: {escape(_safe_text(filters_text or 'Nenhum'))} · {len(rows)} oportunidade(s)",
        styles["BodyText"],
    ))
    story.append(Spacer(1, 4 * mm))
    data = [["Órgão", "Objeto", "Município/UF", "Modalidade", "Fim das propostas", "Valor", "Portal"]]
    for row in rows:
        data.append([
            Paragraph(escape(_safe_text(row.get("agency"))[:55]), styles["BodyText"]),
            Paragraph(escape(_safe_text(row.get("object"))[:210]), styles["BodyText"]),
            f'{row.get("city") or ""}/{row.get("state") or ""}',
            Paragraph(escape(_safe_text(row.get("modality"))[:40]), styles["BodyText"]),
            _format_date(row.get("closing_at")),
            _format_brl(row.get("estimated_value")),
            Paragraph(escape(_safe_text(row.get("source_name") or "Não identificado")[:28]), styles["BodyText"]),
        ])
    story.append(_pdf_table(data, [34*mm, 82*mm, 28*mm, 37*mm, 31*mm, 27*mm, 30*mm]))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Relatório de apoio. Confirme sempre o edital e o portal oficial antes de participar.", styles["BodyText"]))
    doc.build(story)
    return output.getvalue()


def saved_editals_pdf(rows, company_name):
    output = BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10 * mm,
                            leftMargin=10 * mm, topMargin=10 * mm, bottomMargin=10 * mm)
    story = []
    _pdf_header(story, "Meus Editais", company_name, styles)
    story.append(Paragraph(f"Total salvo: {len(rows)} edital(is)", styles["BodyText"]))
    story.append(Spacer(1, 4 * mm))
    data = [["Órgão", "Objeto", "UF", "Status", "Certame", "Valor", "Portal", "Anotações"]]
    for row in rows:
        data.append([
            Paragraph(escape(_safe_text(row.get("agency"))[:45]), styles["BodyText"]),
            Paragraph(escape(_safe_text(row.get("object"))[:170]), styles["BodyText"]),
            row.get("state") or "",
            row.get("stage_display") or "Salvo",
            _format_date(row.get("certame_at")),
            _format_brl(row.get("estimated_value")),
            Paragraph(escape(_safe_text(row.get("source_name") or "Não identificado")[:25]), styles["BodyText"]),
            Paragraph(escape(_safe_text(row.get("notes"))[:140]), styles["BodyText"]),
        ])
    story.append(_pdf_table(data, [31*mm, 68*mm, 10*mm, 22*mm, 29*mm, 25*mm, 28*mm, 55*mm]))
    story.append(Spacer(1, 4 * mm))
    story.append(Paragraph("Relatório de apoio gerencial. Confira as fontes oficiais antes de decidir.", styles["BodyText"]))
    doc.build(story)
    return output.getvalue()
