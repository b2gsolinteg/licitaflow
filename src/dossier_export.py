from __future__ import annotations

from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .journey import journey_progress, list_journey_steps, sync_automatic_completion
from .pricing import item_financials, pricing_summary
from .supplier_directory import list_company_suppliers, list_opportunity_supplier_quotes


NAVY = "152238"
GOLD = "D6A84B"
LIGHT = "F4F6F8"
LOGO_PATH = Path(__file__).resolve().parents[1] / "assets" / "licitanexo-logo.png"
B2G_SIGNATURE = "Desenvolvido por B2G SaaS | Business to Growth"

_STAGE_LABELS = {
    "Nova oportunidade": "Salvo",
    "Em análise": "Analisando",
    "Decisão": "Vou participar",
    "Preparação": "Vou participar",
    "Proposta pronta": "Vou participar",
    "Proposta enviada": "Vou participar",
    "Arquivada": "Encerrado",
    "Ganha": "Encerrado",
    "Perdida": "Encerrado",
}


def _text(value) -> str:
    return "" if value is None else str(value)


def _brl(value) -> str:
    try:
        number = float(value or 0)
    except (TypeError, ValueError):
        number = 0.0
    return f"R$ {number:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _date(value) -> str:
    if not value:
        return ""
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        return str(value).replace("T", " ")[:16]


def _company_info(db, company_id: str) -> dict:
    with db.connect() as conn:
        row = conn.execute("SELECT * FROM companies WHERE id=?", (company_id,)).fetchone()
    return dict(row) if row else {"id": company_id, "name": ""}


def build_opportunity_dossier(db, company_id: str, opportunity: dict) -> dict:
    """Reúne, sob demanda, os dados das áreas do edital sem tornar a navegação mais pesada."""
    opportunity_id = opportunity["id"]
    details = db.get_details(company_id, opportunity_id) or {}
    profile = db.get_company_profile(company_id) or {}
    company = _company_info(db, company_id)

    items = [dict(row) for row in db.list_quote_items(company_id, opportunity_id)]
    quote_groups = list_opportunity_supplier_quotes(db, company_id, opportunity_id)
    supplier_directory = list_company_suppliers(db, company_id, active_only=False)
    directory_by_name = {
        str(row.get("name") or "").strip().casefold(): row for row in supplier_directory
    }

    sync_automatic_completion(db, company_id, opportunity_id)
    journey = list_journey_steps(db, company_id, opportunity_id)
    progress = journey_progress(journey)

    try:
        checklist = [dict(row) for row in db.list_checklist(company_id, opportunity_id)]
    except Exception:
        checklist = []
    try:
        company_documents = [dict(row) for row in db.list_company_documents(company_id)]
    except Exception:
        company_documents = []
    try:
        analyses = [dict(row) for row in db.list_analyses(company_id, opportunity_id)]
    except Exception:
        analyses = []

    findings = []
    for analysis_row in analyses:
        try:
            analysis, analysis_findings = db.get_analysis(company_id, analysis_row["id"])
        except Exception:
            analysis, analysis_findings = None, []
        meta = dict(analysis or analysis_row)
        for finding in analysis_findings or []:
            row = dict(finding)
            row["analysis_id"] = analysis_row.get("id")
            row["analysis_filename"] = meta.get("filename") or analysis_row.get("filename") or ""
            row["analysis_created_at"] = meta.get("created_at") or analysis_row.get("created_at") or ""
            findings.append(row)

    item_rows = []
    quote_rows = []
    for item in items:
        financial = item_financials(item)
        quotes = quote_groups.get(str(item["id"]), [])
        selected = next((row for row in quotes if bool(int(row.get("is_selected") or 0))), None)
        item_rows.append({
            "item_id": item.get("id"),
            "item": item.get("lot_number") or "",
            "descricao": item.get("description") or "",
            "fonte": "PNCP" if str(item.get("source_kind") or "").lower() == "pncp" else "Manual",
            "referencia_fonte": item.get("source_reference") or "",
            "quantidade": financial["quantity"],
            "unidade": item.get("unit_measure") or "",
            "preco_edital": financial["edital_price"],
            "meu_preco": financial["sale_price"],
            "fornecedor": item.get("supplier") or (selected.get("supplier_name") if selected else ""),
            "custo_produto": float(item.get("unit_cost") or 0),
            "frete": float(item.get("freight") or 0),
            "impostos": float(item.get("taxes") or 0),
            "outros_custos": float(item.get("other_costs") or 0) + float(item.get("commission") or 0),
            "custo_final_unitario": financial["cost_unit"],
            "receita_total": financial["revenue"],
            "custo_total": financial["total_cost"],
            "lucro_bruto": financial["profit"],
            "margem_pct": financial["margin_pct"],
            "desconto_referencia_pct": financial["discount_pct"],
        })
        for quote in quotes:
            directory = directory_by_name.get(str(quote.get("supplier_name") or "").strip().casefold()) or {}
            total_unit = sum(float(quote.get(field) or 0) for field in (
                "unit_cost", "freight_unit", "taxes_unit", "other_unit_costs",
            )) + float(item.get("commission") or 0)
            quote_rows.append({
                "item": item.get("lot_number") or "",
                "descricao": item.get("description") or "",
                "fornecedor": quote.get("supplier_name") or "",
                "selecionado": "Sim" if bool(int(quote.get("is_selected") or 0)) else "Não",
                "custo_produto": float(quote.get("unit_cost") or 0),
                "frete": float(quote.get("freight_unit") or 0),
                "impostos": float(quote.get("taxes_unit") or 0),
                "outros": float(quote.get("other_unit_costs") or 0),
                "custo_final_unitario": total_unit,
                "prazo_dias": quote.get("lead_time_days"),
                "observacoes": quote.get("notes") or "",
                "cnpj": directory.get("cnpj") or "",
                "contato": directory.get("contact_name") or "",
                "telefone": directory.get("phone") or "",
                "email": directory.get("email") or "",
                "cidade_uf": "/".join(part for part in (directory.get("city"), directory.get("state")) if part),
                "condicao_pagamento": directory.get("payment_terms") or "",
            })

    summary = pricing_summary(items)
    summary.update({
        "company_name": company.get("name") or profile.get("trade_name") or profile.get("legal_name") or "",
        "company_cnpj": profile.get("cnpj") or company.get("cnpj") or "",
        "agency": opportunity.get("agency") or "",
        "object": opportunity.get("object") or "",
        "city": opportunity.get("city") or "",
        "state": opportunity.get("state") or "",
        "modality": opportunity.get("modality") or "",
        "status": _STAGE_LABELS.get(opportunity.get("stage"), opportunity.get("stage") or ""),
        "pncp_control_number": opportunity.get("pncp_control_number") or "",
        "estimated_value": float(opportunity.get("estimated_value") or 0),
        "published_at": opportunity.get("published_at") or "",
        "opening_at": opportunity.get("opening_at") or "",
        "closing_at": opportunity.get("closing_at") or "",
        "certame_at": details.get("certame_at") or "",
        "source_name": opportunity.get("source_name") or "",
        "source_channel": opportunity.get("source_channel") or "",
        "source_url": opportunity.get("source_url") or "",
        "notes": details.get("notes") or "",
        "journey_done": progress["done"],
        "journey_total": progress["total"],
        "journey_percent": progress["percent"],
        "desired_margin": float(profile.get("desired_margin") or 0),
        "generated_at": datetime.now().isoformat(timespec="seconds"),
    })
    return {
        "summary": summary,
        "items": item_rows,
        "quotes": quote_rows,
        "journey": [dict(row) for row in journey],
        "checklist": checklist,
        "company_documents": company_documents,
        "analyses": analyses,
        "findings": findings,
    }


def _style_excel(writer):
    for sheet in writer.book.worksheets:
        sheet.sheet_view.showGridLines = False
        if sheet.max_row > 1:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[1].height = 28
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for idx, column in enumerate(sheet.columns, 1):
            width = min(max(max((len(str(cell.value or "")) for cell in column), default=8) + 2, 12), 48)
            sheet.column_dimensions[get_column_letter(idx)].width = width


def opportunity_dossier_excel(bundle: dict) -> bytes:
    summary = bundle["summary"]
    summary_rows = [
        ("Empresa", summary.get("company_name")),
        ("CNPJ", summary.get("company_cnpj")),
        ("Órgão", summary.get("agency")),
        ("Objeto", summary.get("object")),
        ("Município/UF", "/".join(part for part in (summary.get("city"), summary.get("state")) if part)),
        ("Modalidade", summary.get("modality")),
        ("Status", summary.get("status")),
        ("Nº PNCP", summary.get("pncp_control_number")),
        ("Publicado em", _date(summary.get("published_at"))),
        ("Fim das propostas", _date(summary.get("closing_at"))),
        ("Data do certame", _date(summary.get("certame_at"))),
        ("Valor estimado do edital", _brl(summary.get("estimated_value"))),
        ("Referência dos itens", _brl(summary.get("edital_total"))),
        ("Minha proposta", _brl(summary.get("revenue"))),
        ("Custo previsto", _brl(summary.get("total_cost"))),
        ("Lucro bruto previsto", _brl(summary.get("profit"))),
        ("Margem geral", f'{float(summary.get("margin_pct") or 0):.2f}%'),
        ("Margem desejada", f'{float(summary.get("desired_margin") or 0):.2f}%'),
        ("Jornada", f'{summary.get("journey_done", 0)}/{summary.get("journey_total", 0)} ({summary.get("journey_percent", 0)}%)'),
        ("Anotações", summary.get("notes")),
        ("Link oficial", summary.get("source_url")),
        ("Gerado em", _date(summary.get("generated_at"))),
    ]

    item_frame = pd.DataFrame(bundle["items"]).rename(columns={
        "item": "Item", "descricao": "Descrição", "fonte": "Fonte", "referencia_fonte": "Referência da fonte",
        "quantidade": "Quantidade", "unidade": "Unidade", "preco_edital": "Preço edital",
        "meu_preco": "Meu preço", "fornecedor": "Fornecedor selecionado", "custo_produto": "Custo produto",
        "frete": "Frete", "impostos": "Impostos", "outros_custos": "Outros custos",
        "custo_final_unitario": "Custo final/un.", "receita_total": "Receita total",
        "custo_total": "Custo total", "lucro_bruto": "Lucro bruto", "margem_pct": "Margem %",
        "desconto_referencia_pct": "Desconto ref. %",
    }).drop(columns=["item_id"], errors="ignore")
    quote_frame = pd.DataFrame(bundle["quotes"]).rename(columns={
        "item": "Item", "descricao": "Descrição", "fornecedor": "Fornecedor", "selecionado": "Selecionado",
        "custo_produto": "Custo produto", "frete": "Frete", "impostos": "Impostos", "outros": "Outros",
        "custo_final_unitario": "Custo final/un.", "prazo_dias": "Prazo (dias)", "observacoes": "Observações",
        "cnpj": "CNPJ", "contato": "Contato", "telefone": "Telefone", "email": "E-mail",
        "cidade_uf": "Cidade/UF", "condicao_pagamento": "Condição de pagamento",
    })
    journey_frame = pd.DataFrame([{
        "Etapa": row.get("title") or "",
        "Concluída": "Sim" if bool(int(row.get("is_done") or 0)) else "Não",
        "Observações": row.get("notes") or "",
    } for row in bundle["journey"]])
    checklist_frame = pd.DataFrame([{
        "Pendência / tarefa": row.get("title") or "",
        "Concluída": "Sim" if bool(int(row.get("is_done") or row.get("done") or 0)) else "Não",
    } for row in bundle["checklist"]])
    documents_frame = pd.DataFrame([{
        "Documento": row.get("document_type") or "",
        "Aplicável": row.get("applicable") or "",
        "Situação": row.get("status") or "",
        "Validade": row.get("expiry_date") or "",
        "Emissor": row.get("issuer") or "",
        "Observações": row.get("notes") or "",
    } for row in bundle["company_documents"]])
    findings_frame = pd.DataFrame([{
        "Análise": row.get("analysis_filename") or "",
        "Data": _date(row.get("analysis_created_at")),
        "Categoria": row.get("category") or "",
        "Severidade": row.get("severity") or "",
        "Título": row.get("title") or "",
        "Página": row.get("page_number") or "",
        "Evidência": row.get("evidence") or "",
        "Ação recomendada": row.get("recommended_action") or "",
    } for row in bundle["findings"]])

    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        pd.DataFrame(summary_rows, columns=["Campo", "Informação"]).to_excel(writer, sheet_name="Resumo", index=False)
        item_frame.to_excel(writer, sheet_name="Itens e preços", index=False)
        quote_frame.to_excel(writer, sheet_name="Cotações fornecedores", index=False)
        journey_frame.to_excel(writer, sheet_name="Jornada", index=False)
        checklist_frame.to_excel(writer, sheet_name="Checklist", index=False)
        documents_frame.to_excel(writer, sheet_name="Documentos empresa", index=False)
        findings_frame.to_excel(writer, sheet_name="Análise edital", index=False)
        _style_excel(writer)
    return output.getvalue()



def _dossier_footer(canvas, doc):
    canvas.saveState()
    page_width, _ = landscape(A4)
    y = 6 * mm
    canvas.setStrokeColor(colors.HexColor("#D6A84B"))
    canvas.setLineWidth(0.7)
    canvas.line(doc.leftMargin, y + 3.5 * mm, page_width - doc.rightMargin, y + 3.5 * mm)
    canvas.setFillColor(colors.HexColor("#667085"))
    canvas.setFont("Helvetica", 7)
    canvas.drawString(doc.leftMargin, y, "LicitaNexo | Desenvolvido por B2G SaaS")
    canvas.drawRightString(page_width - doc.rightMargin, y, f"Página {canvas.getPageNumber()}")
    canvas.restoreState()


def _dossier_brand_story(styles):
    blocks = []
    if LOGO_PATH.exists():
        blocks.append(Image(str(LOGO_PATH), width=58 * mm, height=(58 * mm) / 3.54))
    else:
        blocks.append(Paragraph("<b>LicitaNexo</b>", styles["Title"]))
    blocks.append(Spacer(1, 1.2 * mm))
    blocks.append(Paragraph(
        f'<font color="#152238" size="8"><b>{escape(B2G_SIGNATURE)}</b></font>',
        styles["BodyText"],
    ))
    blocks.append(Spacer(1, 2 * mm))
    rule = Table([[""]], colWidths=[272 * mm], rowHeights=[1.4 * mm])
    rule.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#D6A84B"))]))
    blocks.append(rule)
    blocks.append(Spacer(1, 3 * mm))
    return blocks


def _dossier_why_box(styles):
    text = (
        "Por que preencher tudo: datas organizam prazos; quantidades e preços alimentam a proposta; custos, frete, impostos e fornecedores "
        "mostram a margem real; Jornada, checklist e documentos reduzem pendências. Quanto mais completos e atuais os dados, mais confiável "
        "fica este dossiê para revisão antes da participação."
    )
    box = Table([[Paragraph(escape(text), styles["BodyText"])]], colWidths=[266 * mm])
    box.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F7FA")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#DDE3EA")),
        ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#D6A84B")),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    return box

def _table(data, widths=None, font_size=7):
    table = Table(data, repeatRows=1, colWidths=widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#152238")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#CBD5E1")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), font_size),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F6F8")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def opportunity_dossier_pdf(bundle: dict) -> bytes:
    summary = bundle["summary"]
    styles = getSampleStyleSheet()
    output = BytesIO()
    doc = SimpleDocTemplate(
        output, pagesize=landscape(A4), rightMargin=9 * mm, leftMargin=9 * mm,
        topMargin=10 * mm, bottomMargin=16 * mm,
        title="LicitaNexo - Dossiê de Participação",
        author="B2G SaaS - LicitaNexo",
        subject="Dossiê operacional de participação em licitação",
    )
    story = _dossier_brand_story(styles) + [
        Paragraph("<b>Dossiê de Participação</b>", styles["Title"]),
        Paragraph(
            f'<b>{escape(_text(summary.get("agency")))}</b> · {escape(_text(summary.get("modality")))} · '
            f'{escape(_text(summary.get("city")))}/{escape(_text(summary.get("state")))}',
            styles["Heading2"],
        ),
        Paragraph(escape(_text(summary.get("object"))), styles["BodyText"]),
        Spacer(1, 3 * mm),
    ]

    info = [
        ["Empresa", escape(_text(summary.get("company_name"))), "CNPJ", escape(_text(summary.get("company_cnpj")))],
        ["Nº PNCP", escape(_text(summary.get("pncp_control_number"))), "Status", escape(_text(summary.get("status")))],
        ["Fim das propostas", _date(summary.get("closing_at")), "Certame", _date(summary.get("certame_at"))],
        ["Valor estimado", _brl(summary.get("estimated_value")), "Jornada", f'{summary.get("journey_done", 0)}/{summary.get("journey_total", 0)} ({summary.get("journey_percent", 0)}%)'],
    ]
    story.append(_table(info, [32*mm, 93*mm, 32*mm, 93*mm], 8))
    story.append(Spacer(1, 3 * mm))
    story.append(_dossier_why_box(styles))
    story.append(Spacer(1, 4 * mm))

    financial = [["Itens", "Ref. itens", "Minha proposta", "Custo previsto", "Lucro bruto", "Margem geral"]]
    financial.append([
        str(summary.get("items") or 0), _brl(summary.get("edital_total")), _brl(summary.get("revenue")),
        _brl(summary.get("total_cost")), _brl(summary.get("profit")), f'{float(summary.get("margin_pct") or 0):.2f}%'
    ])
    story.append(Paragraph("<b>Resumo financeiro</b>", styles["Heading2"]))
    story.append(_table(financial, [20*mm, 43*mm, 43*mm, 43*mm, 43*mm, 35*mm], 8))
    story.append(Spacer(1, 4 * mm))

    story.append(Paragraph("<b>Itens, custos, fornecedores e margem</b>", styles["Heading2"]))
    item_table = [["Item", "Descrição", "Fonte", "Qtd./Un.", "Preço edital", "Meu preço", "Fornecedor", "Custo/un.", "Lucro", "Margem"]]
    for row in bundle["items"]:
        item_table.append([
            _text(row.get("item")), Paragraph(escape(_text(row.get("descricao"))[:220]), styles["BodyText"]),
            _text(row.get("fonte")), f'{float(row.get("quantidade") or 0):g} {_text(row.get("unidade"))}'.strip(),
            _brl(row.get("preco_edital")), _brl(row.get("meu_preco")),
            Paragraph(escape(_text(row.get("fornecedor"))[:55]), styles["BodyText"]),
            _brl(row.get("custo_final_unitario")), _brl(row.get("lucro_bruto")),
            f'{float(row.get("margem_pct") or 0):.1f}%',
        ])
    if len(item_table) == 1:
        item_table.append(["—", "Nenhum item cadastrado/importado", "", "", "", "", "", "", "", ""])
    story.append(_table(item_table, [11*mm, 65*mm, 16*mm, 24*mm, 28*mm, 28*mm, 38*mm, 28*mm, 28*mm, 19*mm], 6.5))
    story.append(PageBreak())

    story.append(Paragraph("<b>Jornada e checklist</b>", styles["Heading2"]))
    journey_table = [["Situação", "Etapa", "Observações"]]
    for row in bundle["journey"]:
        journey_table.append([
            "Concluída" if bool(int(row.get("is_done") or 0)) else "Pendente",
            Paragraph(escape(_text(row.get("title"))), styles["BodyText"]),
            Paragraph(escape(_text(row.get("notes"))), styles["BodyText"]),
        ])
    story.append(_table(journey_table, [28*mm, 115*mm, 112*mm], 7))
    if bundle["checklist"]:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("<b>Pendências detectadas / checklist</b>", styles["Heading3"]))
        for row in bundle["checklist"]:
            done = bool(int(row.get("is_done") or row.get("done") or 0))
            story.append(Paragraph(f'{"OK" if done else "PENDENTE"} · {escape(_text(row.get("title")))}', styles["BodyText"]))

    if bundle["company_documents"]:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("<b>Documentos da empresa</b>", styles["Heading2"]))
        docs = [["Documento", "Aplicável", "Situação", "Validade", "Observações"]]
        for row in bundle["company_documents"]:
            docs.append([
                Paragraph(escape(_text(row.get("document_type"))), styles["BodyText"]),
                _text(row.get("applicable")), _text(row.get("status")), _text(row.get("expiry_date")),
                Paragraph(escape(_text(row.get("notes"))[:160]), styles["BodyText"]),
            ])
        story.append(_table(docs, [62*mm, 26*mm, 31*mm, 29*mm, 105*mm], 7))

    if bundle["findings"]:
        story.append(PageBreak())
        story.append(Paragraph("<b>Análise do edital e anexos</b>", styles["Heading2"]))
        for row in bundle["findings"][-40:]:
            page = f' · pág. {row.get("page_number")}' if row.get("page_number") else ""
            story.append(Paragraph(
                f'<b>{escape(_text(row.get("category")))} · {escape(_text(row.get("title")))}</b>{page}',
                styles["BodyText"],
            ))
            if row.get("evidence"):
                story.append(Paragraph(escape(_text(row.get("evidence"))[:900]), styles["BodyText"]))
            if row.get("recommended_action"):
                story.append(Paragraph(f'Ação: {escape(_text(row.get("recommended_action"))[:500])}', styles["BodyText"]))
            story.append(Spacer(1, 2 * mm))

    if summary.get("notes"):
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("<b>Minhas anotações</b>", styles["Heading2"]))
        story.append(Paragraph(escape(_text(summary.get("notes"))), styles["BodyText"]))

    story.append(Spacer(1, 5 * mm))
    story.append(Paragraph(
        "Documento operacional de apoio gerado pelo LicitaNexo. Valores identificados como PNCP vêm da consulta estruturada oficial; "
        "itens manuais dependem da conferência do usuário. Confirme sempre edital, Termo de Referência, anexos e portal oficial antes do envio da proposta.",
        styles["BodyText"],
    ))
    doc.build(story, onFirstPage=_dossier_footer, onLaterPages=_dossier_footer)
    return output.getvalue()
