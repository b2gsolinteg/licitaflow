import io
import sys
import unittest
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from src.dossier_export import opportunity_dossier_excel, opportunity_dossier_pdf


class DossierExportTests(unittest.TestCase):
    def sample_bundle(self):
        return {
            "summary": {
                "company_name": "Empresa Teste",
                "company_cnpj": "12.345.678/0001-90",
                "agency": "Prefeitura Teste",
                "object": "Aquisição de papel A4",
                "city": "São Paulo",
                "state": "SP",
                "modality": "Pregão eletrônico",
                "status": "Vou participar",
                "pncp_control_number": "12345678000190-1-1/2026",
                "published_at": "2026-08-01T10:00:00",
                "closing_at": "2026-08-20T09:00:00",
                "certame_at": "2026-08-20T10:00:00",
                "estimated_value": 3200,
                "items": 1,
                "edital_total": 3200,
                "revenue": 2790,
                "total_cost": 2185,
                "profit": 605,
                "margin_pct": 21.6845,
                "desired_margin": 15,
                "journey_done": 5,
                "journey_total": 8,
                "journey_percent": 63,
                "notes": "Conferir amostra antes da sessão.",
                "source_url": "https://example.invalid/edital",
                "generated_at": "2026-08-17T12:30:00",
            },
            "items": [{
                "item_id": "i1", "item": "1", "descricao": "Papel A4", "fonte": "PNCP",
                "referencia_fonte": "pncp:controle:1", "quantidade": 100, "unidade": "RESMA",
                "preco_edital": 32, "meu_preco": 27.9, "fornecedor": "Fornecedor A",
                "custo_produto": 20.85, "frete": 0.5, "impostos": 0.5, "outros_custos": 0,
                "custo_final_unitario": 21.85, "receita_total": 2790, "custo_total": 2185,
                "lucro_bruto": 605, "margem_pct": 21.6845, "desconto_referencia_pct": 12.8125,
            }],
            "quotes": [{
                "item": "1", "descricao": "Papel A4", "fornecedor": "Fornecedor A", "selecionado": "Sim",
                "custo_produto": 20.85, "frete": 0.5, "impostos": 0.5, "outros": 0,
                "custo_final_unitario": 21.85, "prazo_dias": 5, "observacoes": "À vista",
                "cnpj": "11.111.111/0001-11", "contato": "Ana", "telefone": "11999999999",
                "email": "ana@example.com", "cidade_uf": "São Paulo/SP", "condicao_pagamento": "28 dias",
            }],
            "journey": [{"title": "Analisar edital e anexos", "is_done": 1, "notes": ""}],
            "checklist": [{"title": "Conferir CNDT", "is_done": 0}],
            "company_documents": [{
                "document_type": "CNDT", "applicable": "Sim", "status": "Válido",
                "expiry_date": "2026-10-01", "issuer": "TST", "notes": "",
            }],
            "analyses": [{"id": "a1", "filename": "edital.pdf", "created_at": "2026-08-17T11:00:00"}],
            "findings": [{
                "analysis_filename": "edital.pdf", "analysis_created_at": "2026-08-17T11:00:00",
                "category": "Documentação", "severity": "Amarelo", "title": "Regularidade fiscal",
                "page_number": 12, "evidence": "Certidões válidas na data da sessão.",
                "recommended_action": "Conferir validade.",
            }],
        }

    def test_excel_has_all_operational_sheets(self):
        payload = opportunity_dossier_excel(self.sample_bundle())
        workbook = load_workbook(io.BytesIO(payload), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            ["Resumo", "Itens e preços", "Cotações fornecedores", "Jornada", "Checklist", "Documentos empresa", "Análise edital"],
        )
        self.assertEqual(workbook["Itens e preços"]["C2"].value, "PNCP")

    def test_pdf_is_generated(self):
        payload = opportunity_dossier_pdf(self.sample_bundle())
        self.assertTrue(payload.startswith(b"%PDF"))
        self.assertGreater(len(payload), 1500)
        reader = PdfReader(io.BytesIO(payload))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        self.assertIn("Desenvolvido por B2G SaaS", text)
        self.assertIn("Por que preencher tudo", text)


if __name__ == "__main__":
    unittest.main()
